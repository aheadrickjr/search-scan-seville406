"""Builds the Excel evidence register workbook for a single scan.

Sheets (Phase 1):
  1. Executive Summary
  2. Current Results
  3. PIR-Associated Listings
  4. Franke Listings
  5. Link Validation
  6. Manual Review

Change-since-last-scan comparison and an Action Register are Phase 3 --
not built here, but the append-only DB schema this reads from supports it.
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import Classification, Scan, ScanQuery, SearchResult, UrlCheck

PIR_ASSOCIATED_CLASSIFICATIONS = {
    "Active PIR listing",
    "Active third-party listing attributed to PIR",
    "Dead PIR page",
    "Redirect to PIR inventory",
    "Redirect to unrelated property",
}

MANUAL_REVIEW_CLASSIFICATIONS = {
    "Blocked or inaccessible",
    "Unclear — manual review required",
}


def _load_scan_dataframe(session: Session, scan_id: int) -> pd.DataFrame:
    search_result_rows = session.execute(
        select(SearchResult, ScanQuery.query_text)
        .join(ScanQuery, SearchResult.scan_query_id == ScanQuery.id)
        .where(SearchResult.scan_id == scan_id)
    ).all()

    by_url: dict[str, dict] = {}
    for result, query_text in search_result_rows:
        entry = by_url.setdefault(
            result.normalized_url,
            {
                "normalized_url": result.normalized_url,
                "domain": result.domain,
                "best_position": result.position,
                "title": result.title,
                "snippet": result.snippet,
                "queries": set(),
            },
        )
        entry["queries"].add(query_text)
        if result.position < entry["best_position"]:
            entry["best_position"] = result.position
            entry["title"] = result.title
            entry["snippet"] = result.snippet

    url_checks = {
        row.normalized_url: row
        for row in session.execute(
            select(UrlCheck).where(UrlCheck.scan_id == scan_id)
        ).scalars()
    }
    classifications = {
        row.normalized_url: row
        for row in session.execute(
            select(Classification).where(Classification.scan_id == scan_id)
        ).scalars()
    }

    records = []
    for normalized_url, entry in by_url.items():
        check = url_checks.get(normalized_url)
        clas = classifications.get(normalized_url)
        records.append(
            {
                "Normalized URL": normalized_url,
                "Domain": entry["domain"],
                "Best Rank Position": entry["best_position"],
                "Matched Queries": "; ".join(sorted(entry["queries"])),
                "Title": entry["title"],
                "Snippet": entry["snippet"],
                "Status Code": check.status_code if check else None,
                "Redirect Chain": (
                    " -> ".join(json.loads(check.redirect_chain_json))
                    if check and check.redirect_chain_json
                    else ""
                ),
                "Final URL": check.final_url if check else None,
                "Response Time (ms)": check.response_time_ms if check else None,
                "Accessible": check.accessible if check else None,
                "Checked At (UTC)": check.checked_at if check else None,
                "Robots.txt Disallowed Content Check": (
                    check.robots_disallowed if check else None
                ),
                "Classification": clas.classification if clas else "Unclear — manual review required",
                "Confidence": clas.confidence if clas else "Low",
                "Mentions Seville 406": clas.mentions_seville_406 if clas else False,
                "Mentions PIR": clas.mentions_pir if clas else False,
                "Mentions Franke": clas.mentions_franke if clas else False,
                "Mentions Only In Hidden Metadata": (
                    clas.mentions_only_in_hidden_metadata if clas else False
                ),
                "Booking Language Detected": (
                    clas.booking_language_detected if clas else False
                ),
                "Booking Status": (
                    clas.booking_status_text if clas else "Unable to determine"
                ),
                "Manual Review Flag": clas.manual_review_flag if clas else True,
                "Notes": clas.notes if clas else "",
            }
        )

    df = pd.DataFrame.from_records(records)
    if not df.empty:
        df = df.sort_values("Best Rank Position").reset_index(drop=True)
    return df


def _autosize_and_bold_header(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    worksheet = writer.sheets[sheet_name]
    for col_idx, column in enumerate(df.columns, start=1):
        max_len = max([len(str(column))] + [len(str(v)) for v in df[column].astype(str).tolist()[:200]])
        worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)
        worksheet.cell(row=1, column=col_idx).font = Font(bold=True)
    worksheet.freeze_panes = "A2"


def export_scan_to_excel(session: Session, scan_id: int, output_path: Path) -> Path:
    scan = session.get(Scan, scan_id)
    if scan is None:
        raise ValueError(f"No scan with id {scan_id}")

    df = _load_scan_dataframe(session, scan_id)

    pir_df = df[df["Classification"].isin(PIR_ASSOCIATED_CLASSIFICATIONS)] if not df.empty else df
    franke_df = df[df["Classification"] == "Franke Rentals listing"] if not df.empty else df
    manual_review_df = (
        df[df["Manual Review Flag"] | df["Classification"].isin(MANUAL_REVIEW_CLASSIFICATIONS)]
        if not df.empty
        else df
    )
    link_validation_columns = [
        "Normalized URL",
        "Domain",
        "Status Code",
        "Redirect Chain",
        "Final URL",
        "Response Time (ms)",
        "Accessible",
        "Checked At (UTC)",
        "Robots.txt Disallowed Content Check",
        "Classification",
    ]
    link_validation_df = df[link_validation_columns] if not df.empty else df

    summary_rows = [
        ("Scan date (UTC)", scan.run_at),
        ("Search provider", scan.provider),
        ("Total unique results", len(df)),
        ("Total PIR-associated results", len(pir_df)),
        (
            "Active PIR listings",
            int((df["Classification"] == "Active PIR listing").sum()) if not df.empty else 0,
        ),
        (
            "Dead PIR pages",
            int((df["Classification"] == "Dead PIR page").sum()) if not df.empty else 0,
        ),
        (
            "Redirects (PIR-related)",
            int(
                df["Classification"]
                .isin(["Redirect to PIR inventory", "Redirect to unrelated property"])
                .sum()
            )
            if not df.empty
            else 0,
        ),
        (
            "Results with booking controls detected",
            int(df["Booking Language Detected"].sum()) if not df.empty else 0,
        ),
        ("Franke Rentals listings found", len(franke_df)),
        ("Flagged for manual review", len(manual_review_df)),
    ]
    summary_df = pd.DataFrame(summary_rows, columns=["Metric", "Value"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        sheets = {
            "Executive Summary": summary_df,
            "Current Results": df,
            "PIR-Associated Listings": pir_df,
            "Franke Listings": franke_df,
            "Link Validation": link_validation_df,
            "Manual Review": manual_review_df,
        }
        for sheet_name, sheet_df in sheets.items():
            sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)
            _autosize_and_bold_header(writer, sheet_name, sheet_df)

    return output_path
